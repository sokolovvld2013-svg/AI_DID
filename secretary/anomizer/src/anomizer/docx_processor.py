from __future__ import annotations
import re
from dataclasses import dataclass
from typing import List, Dict, Optional, Tuple
from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.table import Table
from docx.oxml.ns import qn

from .registry import EntityRegistry, Entity
from .ner import ExtractedEntity, NERExtractor


@dataclass
class TextSegment:
    text: str
    runs: List[Run]
    start_offset: int
    end_offset: int


def _iter_paragraph_runs(paragraph: Paragraph):
    """Yield all runs including those inside hyperlinks."""
    for run in paragraph.runs:
        yield run
    # Check for hyperlinks
    for hyperlink in paragraph._p.findall(qn('w:hyperlink')):
        for run in hyperlink.findall(qn('w:r')):
            # Create a Run-like object from the XML element
            from docx.text.run import Run as RunClass
            yield RunClass(run, paragraph)


class DocxProcessor:
    def __init__(self):
        self.ner_extractor = NERExtractor()
        self.registry = EntityRegistry()

    def process(self, input_path: str, output_path: str, extra_orgs: Optional[List[str]] = None) -> Dict[str, Entity]:
        doc = Document(input_path)
        
        full_text, segments = self._extract_text_with_segments(doc)
        
        entities = self.ner_extractor.extract(full_text, extra_orgs)
        
        for entity in entities:
            self.registry.register(Entity(
                text=entity.text,
                normalized=entity.normalized,
                entity_type=entity.entity_type,
            ), entity.start, entity.end)
        
        self._replace_in_segments(segments)
        
        doc.save(output_path)
        
        return {e.placeholder: e for e in self.registry.get_all_entities()}

    def _extract_text_with_segments(self, doc: Document) -> Tuple[str, List[TextSegment]]:
        segments = []
        full_text_parts = []
        current_offset = 0
        
        for paragraph in doc.paragraphs:
            seg = self._process_paragraph(paragraph, current_offset)
            if seg:
                segments.append(seg)
                full_text_parts.append(seg.text)
                current_offset = seg.end_offset
            full_text_parts.append('\n')
            current_offset += 1
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        seg = self._process_paragraph(paragraph, current_offset)
                        if seg:
                            segments.append(seg)
                            full_text_parts.append(seg.text)
                            current_offset = seg.end_offset
                        full_text_parts.append('\n')
                        current_offset += 1
        
        full_text = ''.join(full_text_parts)
        return full_text, segments

    def _process_paragraph(self, paragraph: Paragraph, base_offset: int) -> Optional[TextSegment]:
        runs_with_text = []
        
        for run in _iter_paragraph_runs(paragraph):
            text = run.text
            if text:
                runs_with_text.append((run, text))
        
        if not runs_with_text:
            return None
        
        combined_text = ''.join(text for _, text in runs_with_text)
        return TextSegment(
            text=combined_text,
            runs=[run for run, _ in runs_with_text],
            start_offset=base_offset,
            end_offset=base_offset + len(combined_text)
        )

    def _replace_in_segments(self, segments: List[TextSegment]):
        # Собираем планы замен на каждого сегмент в координатах ОРИГИНАЛЬНОГО текста сегмента,
        # затем применяем сегмент целиком за один проход (пересечения: родитель выигрывает).
        plans_per_segment: List[List[Tuple[int, int, str]]] = [[] for _ in segments]

        for entity in self.registry.get_all_entities():
            for start, end in entity.positions:
                for segment in segments:
                    if end <= segment.start_offset or start >= segment.end_offset:
                        continue
                    local_start = max(0, start - segment.start_offset)
                    local_end = min(len(segment.text), end - segment.start_offset)
                    if local_start < local_end:
                        plans_per_segment[segments.index(segment)].append(
                            (local_start, local_end, entity.placeholder)
                        )

        for segment, plans in zip(segments, plans_per_segment):
            if plans:
                self._apply_segment_replacements(segment, plans)

    def _apply_segment_replacements(self, segment: TextSegment, plans: List[Tuple[int, int, str]]):
        if not segment.runs:
            return

        seg_text = segment.text
        n = len(seg_text)
        replaced: List[Optional[str]] = [None] * n

        # Сначала родительские (большие) спаны — они выигрывают у вложенных
        plans.sort(key=lambda p: (p[0], -p[1]))
        for ls, le, ph in plans:
            if any(x is not None for x in replaced[ls:le]):
                continue
            for i in range(ls, le):
                replaced[i] = ph

        # Финальный текст сегмента
        pieces = []  # (orig_start, orig_end, text_or_placeholder)
        i = 0
        while i < n:
            if replaced[i] is not None:
                ph = replaced[i]
                j = i
                while j < n and replaced[j] == ph:
                    j += 1
                pieces.append((i, j, ph))
                i = j
            else:
                pieces.append((i, i + 1, seg_text[i]))
                i += 1

        # Раскладываем обратно по run'ам (по границам ОРИГИНАЛЬНЫХ текстов run'ов)
        run_offsets = []
        cur = 0
        for run in segment.runs:
            run_offsets.append((cur, cur + len(run.text), run))
            cur += len(run.text)

        for rs, re, run in run_offsets:
            parts = []
            for ps, pe, payload in pieces:
                if ps < rs or ps >= re:
                    continue
                # Замена выводится только в том run, где она начинается;
                # обычные символы попали в куски по одному, поэтому ps == их индексу
                parts.append(payload)
            run.text = ''.join(parts)


def process_docx(input_path: str, output_path: str, extra_orgs: Optional[List[str]] = None) -> Dict[str, Entity]:
    processor = DocxProcessor()
    return processor.process(input_path, output_path, extra_orgs)