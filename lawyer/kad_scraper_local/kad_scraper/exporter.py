"""Выгрузка в Excel и промежуточные чекпойнты.

Основной формат — ``.xlsx`` через pandas + openpyxl. Файл формируется в памяти
(``BytesIO``), поэтому его можно отдать в ``st.download_button`` без
промежуточного диска.

Чекпойнты нужны потому, что обход 30-50 страниц с паузами по 5-10 с занимает
10-30 минут: если процесс упадёт на 27-й странице, уже собранное не должно
потеряться.
"""

from __future__ import annotations

import io
import json
import logging
import re
import time
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .models import COLUMNS, CaseRecord, merge_parties

log = logging.getLogger("kad.exporter")

#: Поля со списками сторон — по ним нужен слияние, а не «непустое положить».
_PARTY_FIELDS = ("plaintiffs", "respondents", "third_parties", "other_persons")


def _merge_saved(
    old: dict[str, Any], new: dict[str, Any]
) -> dict[str, Any]:
    """Дополняет сохранённую запись новой, не теряя заполненные поля.

    Источники данных о деле неравноценны: таблица результатов знает ИНН, но
    не показывает третьих лиц, а карточка — наоборот. Поэтому побеждает
    непустое значение из новой записи, а всё, чего в ней нет, остаётся на
    месте. Пустая строка никогда не затирает заполненное значение.
    """
    merged = dict(old)
    for field_name, value in new.items():
        if field_name in _PARTY_FIELDS:
            if value:
                merged[field_name] = merge_parties(
                    old.get(field_name) or [], value
                )
            continue
        if field_name == "enriched":
            merged[field_name] = bool(old.get(field_name)) or bool(value)
            continue
        if value not in (None, "", 0) or not merged.get(field_name):
            merged[field_name] = value
    return merged

#: Цветовая схема листа.
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_LINK_FONT = Font(color="0563C1", underline="single")

#: Колонки, в которых длинный текст переносится по словам.
_WRAP_HEADERS = frozenset(
    {"Истцы", "Ответчики", "Третьи лица", "Иные лица", "Статус", "Участники (адреса)"}
)

#: Ширина колонок в символах, задаётся явно — иначе Excel обрезает длинные
#: ФИО участников дела.
_COLUMN_WIDTHS: dict[str, int] = {
    "Источник информации": 22,
    "Ссылка на источник": 46,
    "Номер дела": 18,
    "Истцы": 52,
    "Ответчики": 52,
    "Третьи лица": 40,
    "Иные лица": 40,
    "Номер дела по инстанции": 26,
    "Сумма иска": 22,
    "Статус": 44,
    "Дата регистрации": 16,
    "Категория дела": 18,
    "Судья": 24,
    "Суд": 34,
    "ИНН истцов": 18,
    "ИНН ответчиков": 18,
    "Участники (адреса)": 60,
    "Страница КАД": 12,
}


# --------------------------------------------------------------------------- #
# DataFrame
# --------------------------------------------------------------------------- #


def records_to_dataframe(records: Sequence[CaseRecord]) -> pd.DataFrame:
    """Превращает записи в DataFrame с фиксированным порядком колонок."""
    rows = [record.as_row() for record in records]
    frame = pd.DataFrame(rows, columns=list(COLUMNS))
    if frame.empty:
        return frame
    # Ставим пустые значения вместо NaN — в Excel это выглядит аккуратнее
    # и не превращается в текст «nan».
    return frame.fillna("")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #


def records_to_excel(
    records: Sequence[CaseRecord],
    *,
    inn: str | None = None,
    sheet_title: str | None = None,
) -> bytes:
    """Формирует ``.xlsx`` в памяти и возвращает его байты."""
    frame = records_to_dataframe(records)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(
            writer,
            index=False,
            sheet_name=sheet_title or _safe_sheet_name(f"Дела по ИНН {inn}" if inn else "Дела"),
        )
        data_sheet = next(iter(writer.sheets.values()))
        _style_sheet(data_sheet)
        _fill_meta(writer.book.create_sheet("Источник"), records, inn)

    return buffer.getvalue()


def _style_sheet(worksheet: Any) -> None:
    """Оформление листа: шапка, автофильтр, закреплённая шапка, ширины, ссылки."""
    if worksheet is None or worksheet.max_row < 1:
        return

    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 30

    last_col = worksheet.max_column
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(last_col)}{max(worksheet.max_row, 1)}"
    )
    worksheet.freeze_panes = "A2"

    headers = {cell.value: cell.column for cell in worksheet[1]}
    for header, column in headers.items():
        width = _COLUMN_WIDTHS.get(str(header), 22)
        worksheet.column_dimensions[get_column_letter(column)].width = width

    # Кликабельная ссылка на карточку дела + перенос строк в длинных колонках.
    link_column = headers.get("Ссылка на источник")
    wrap_columns = {
        column for header, column in headers.items() if str(header) in _WRAP_HEADERS
    }

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top", wrap_text=cell.column in wrap_columns
            )
        if link_column:
            cell = row[link_column - 1]
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = _LINK_FONT


def _fill_meta(worksheet: Any, records: Sequence[CaseRecord], inn: str | None) -> None:
    """Лист «Источник» — что, когда и откуда выгружено."""
    pages = sorted({record.kad_page for record in records if record.kad_page})
    rows: list[tuple[str, Any]] = [
        ("Источник информации", config.SOURCE_INFO),
        ("Название источника", "Картотека арбитражных дел"),
        ("ИНН", inn or ""),
        ("Записей выгружено", len(records)),
        ("Страниц КАД обработано", f"{min(pages)}–{max(pages)}" if pages else "—"),
        ("Уникальных номеров дел", len({r.case_number for r in records if r.case_number})),
        ("Обогащение карточек", "да" if any(r.enriched for r in records) else "нет"),
        ("Дата выгрузки", time.strftime("%d.%m.%Y %H:%M:%S")),
    ]

    for row_index, (key, value) in enumerate(rows, start=1):
        worksheet.cell(row=row_index, column=1, value=key).font = Font(bold=True)
        worksheet.cell(row=row_index, column=2, value=value)

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 62
    if len(records) and records[0].url:
        cell = worksheet.cell(row=1, column=2)
        cell.hyperlink = records[0].url
        cell.font = _LINK_FONT


def _safe_sheet_name(name: str) -> str:
    """Имя листа в Excel: не длиннее 31 символа и без запрещённых знаков."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", name).strip("' ")
    return (cleaned or "Дела")[:31]


def suggest_filename(inn: str, *, stamp: str | None = None) -> str:
    """Имя файла выгрузки: ``kad_arbitr_INN_2026-10-03_12-30-00.xlsx``."""
    stamp = stamp or time.strftime("%Y-%m-%d_%H-%M-%S")
    safe_inn = re.sub(r"[^0-9A-Za-z]+", "_", str(inn)) or "inn"
    return f"kad_arbitr_{safe_inn}_{stamp}.xlsx"


# --------------------------------------------------------------------------- #
# Чекпойнты
# --------------------------------------------------------------------------- #


class CheckpointStore:
    """Промежуточное сохранение собранных дел в ``output/``.

    Пишутся два файла:

    * ``kad_checkpoint_<ИНН>.csv`` — плоские строки листа Excel. Их можно
      открыть и посмотреть без приложения.
    * ``kad_checkpoint_<ИНН>.json`` — полные структурированные записи плюс
      метаданные обхода.

    Продолжение сбора идёт именно по JSON: в CSV стороны склеены в одну строку
    через ``; ``, и разбирать её обратно — потеря данных (ИНН и адреса теряются).
    """

    def __init__(self, directory: Path | str = "output", inn: str | None = None) -> None:
        self.directory = Path(directory)
        self.inn = str(inn or "").strip() or "inn"
        self.directory.mkdir(parents=True, exist_ok=True)

    @property
    def csv_path(self) -> Path:
        return self.directory / f"kad_checkpoint_{_safe(self.inn)}.csv"

    @property
    def meta_path(self) -> Path:
        return self.directory / f"kad_checkpoint_{_safe(self.inn)}.json"

    def load(self) -> tuple[list[CaseRecord], dict[str, Any]]:
        """Читает чекпойнт. Возвращает (``записи``, ``метаданные``)."""
        if not self.meta_path.exists():
            return [], {}
        try:
            payload = json.loads(self.meta_path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            log.warning("Чекпойнт %s не читается: %s", self.meta_path, exc)
            return [], {}

        if not isinstance(payload, dict):
            return [], {}
        meta = payload.get("meta") if isinstance(payload.get("meta"), dict) else {}
        raw_records = payload.get("records")
        if not isinstance(raw_records, list):
            return [], meta

        records = [
            record
            for raw in raw_records
            if isinstance(raw, dict) and (record := CaseRecord.from_saved(raw))
        ]
        log.info("Чекпойнт %s: восстановлено записей — %d", self.inn, len(records))
        return records, meta

    def save(
        self,
        records: Sequence[CaseRecord],
        meta: dict[str, Any] | None = None,
        *,
        merge: bool = True,
    ) -> int:
        """Сохраняет прогресс: сначала JSON, затем CSV через rename.

        По умолчанию чекпойнт **дополняется**, а не перезаписывается. Иначе
        пробный запуск на 10 строк с лимитом стирал бы уже собранные 787 дел,
        и восстановить их было бы нечем. Чтобы начать с нуля, чекпойнт нужно
        удалить явно через :meth:`clear`.

        Возвращает итоговое число записей в чекпойнте.
        """
        if not records:
            return 0

        payload_records = [record.as_dict() for record in records]
        if merge:
            payload_records = self._merge_with_disk(payload_records)

        payload = {
            "inn": self.inn,
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "count": len(payload_records),
            "meta": meta or {},
            "records": payload_records,
        }
        json_tmp = self.meta_path.with_suffix(".json.tmp")
        try:
            json_tmp.write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )
            json_tmp.replace(self.meta_path)
        except OSError as exc:
            log.warning("Не удалось сохранить чекпойнт: %s", exc)
            json_tmp.unlink(missing_ok=True)
            return 0

        # CSV — вспомогательный, удобный для человека. Если он не записался,
        # это не повод терять JSON с полными данными.
        try:
            merged = [
                record
                for raw in payload_records
                if (record := CaseRecord.from_saved(raw))
            ]
            frame = records_to_dataframe(merged)
            csv_tmp = self.csv_path.with_suffix(".csv.tmp")
            frame.to_csv(csv_tmp, index=False)
            csv_tmp.replace(self.csv_path)
        except (OSError, ValueError) as exc:
            log.warning("Не удалось сохранить CSV чекпойнта: %s", exc)

        return len(payload_records)

    def _merge_with_disk(self, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """Дополняет ``records`` тем, что уже лежит в чекпойнте на диске.

        Записи сопоставляются по :attr:`CaseRecord.dedup_key`. Всё, чего нет
        в новых данных, берётся из прежних, поэтому частичный прогон не
        уменьшает выгрузку.
        """
        try:
            payload = json.loads(self.meta_path.read_text(encoding="utf-8"))
            stored = payload.get("records") if isinstance(payload, dict) else None
        except (OSError, ValueError):
            stored = None

        if not isinstance(stored, list) or not stored:
            return records

        merged: dict[str, dict[str, Any]] = {}
        order: list[str] = []

        for raw in stored:
            if not isinstance(raw, dict):
                continue
            record = CaseRecord.from_saved(raw)
            if record is None:
                continue
            key = record.dedup_key
            if key not in merged:
                order.append(key)
            merged[key] = record.as_dict()

        added = 0
        for raw in records:
            record = CaseRecord.from_saved(raw)
            if record is None:
                continue
            key = record.dedup_key
            if key in merged:
                merged[key] = _merge_saved(merged[key], raw)
            else:
                order.append(key)
                merged[key] = raw
                added += 1

        kept = len(order)
        if added or kept != len(records):
            log.info(
                "Чекпойнт %s: сохранено записей — %d (добавлено %d)",
                self.inn,
                kept,
                added,
            )
        return [merged[key] for key in order]

    def clear(self) -> None:
        for path in (self.csv_path, self.meta_path):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass


def _safe(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z]+", "_", value) or "inn"
